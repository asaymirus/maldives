"""Phase G: build output/documents.csv|json, output/coverage.xlsx (with live formulas), pages index, REPORT.md."""
import os, re, json, csv, datetime, collections
from urllib.parse import urlparse
import lib
from lib import load_docs, mark_latest, save_docs, OUT, ROOT
from resorts import RESORTS, BY_ID
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DOC_TYPES = ["factsheet", "wedding", "events", "map", "dive_map", "spa_menu", "dining_menu", "dive_prices", "excursions",
             "kids", "villa_plans", "all_inclusive", "sustainability", "calendar", "brochure", "press_kit", "other", "obsolete"]
DOC_FIELDS = ["resort_id", "resort_name", "doc_type", "file_name", "edition_year", "valid_from", "valid_to", "pages", "size_kb",
              "text_layer", "source", "source_type", "url", "all_urls", "sha256", "local_path", "r2_key", "status", "is_latest",
              "match_status", "checked_on", "notes"]

RENAMES_KNOWN = [
    (5, "Bolifushi Island Resort", "OZEN Reserve Bolifushi (confirm)"), (22, "Barceló Whale Lagoon Maldives", "Meliá Whale Lagoon Maldives"),
    (25, "Raffles Maldives Meradhoo Resort", "Halcyon Maldives"), (33, "Dhigali Maldives", "NIVA Dhigali"),
    (39, "NH Collection Maldives Havodda", "Amari Havodda (older name)"), (41, "Ozen by Atmosphere at Maadhoo", "OZEN Life Maadhoo"),
    (50, "Sun Siyam Olhuveli", "Olhuveli Beach & Spa, file code SSO"), (58, "Sirru Fen Fushi", "Fairmont Maldives Sirru Fen Fushi"),
    (85, "Centara Mirage Lagoon & Centara Grand Lagoon", "two resorts sharing one entry"), (92, "Amaya Kuda Rah", "NH Maldives Kuda Rah"),
    (96, "Dhawa Ihuru", "Angsana Ihuru (older name), file code DHMVIH"), (119, "Outrigger Maldives Maafushivaru", "Lti Maafushivaru (older name)"),
    (133, "Oblu by Atmosphere at Helengeli", "OBLU Nature Helengeli by Sentido"), (148, "Sun Siyam Iru Fushi", "The Sun Siyam Iru Fushi, file code TSSI"),
    (155, "Constance Halaveli", "file code CHRG"), (165, "Velassaru Maldives", "NIVA Velassaru"), (168, "You & Me Maldives", "You & Me by Cocoon"),
    (175, "Villa Nautica Paradise Island", "Paradise Island Resort"), (29, "Dheruhfinolhu / Mabinhura", "Jawakara Islands (shared documents)"),
    (83, "Dheruhfinolhu / Mabinhura", "Jawakara Islands (shared documents)"), (38, "Veligandu Maldives Resort Island", "'Veli' in Crown & Champa filenames"),
]


def doc_rows(docs):
    rows = []
    for d in sorted(docs.values(), key=lambda d: (d["resort_id"], DOC_TYPES.index(d["doc_type"]) if d["doc_type"] in DOC_TYPES else 99, -(d["edition_year"] or 0))):
        for rid in d.get("resort_ids") or [d["resort_id"]]:
            row = {k: d.get(k) for k in DOC_FIELDS}
            row["resort_id"] = rid
            row["resort_name"] = BY_ID[rid]["name"]
            row["all_urls"] = " | ".join(d.get("all_urls", []))
            rows.append(row)
    return rows


def write_documents(docs):
    rows = doc_rows(docs)
    with open(os.path.join(OUT, "documents.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=DOC_FIELDS)
        w.writeheader()
        w.writerows(rows)
    with open(os.path.join(OUT, "documents.json"), "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=1)
    return rows


ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean(v):
    if isinstance(v, str):
        v = ILLEGAL.sub("", v)
        if v.startswith(("=", "+", "-", "@")) and not v.startswith(("=IFERROR", "=COUNTIFS", "=HYPERLINK", "=MAXIFS")):
            v = "'" + v
    return v


def append(ws, row):
    ws.append([clean(c) for c in row])


def style_header(ws):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"


def autosize(ws, maxw=60):
    for i, col in enumerate(ws.columns, 1):
        w = max((len(str(c.value)) if c.value is not None else 0) for c in list(col)[:200])
        ws.column_dimensions[get_column_letter(i)].width = min(maxw, max(8, w + 2))


def build_workbook(docs, rows):
    wb = Workbook()
    # ---- Documents sheet (source for formulas)
    wsd = wb.active
    wsd.title = "Documents"
    append(wsd, DOC_FIELDS)
    for r in rows:
        append(wsd, [r.get(k) if not isinstance(r.get(k), bool) else ("Y" if r.get(k) else "N") for k in DOC_FIELDS])
    style_header(wsd)
    autosize(wsd, 50)
    nd = len(rows) + 1
    col = {k: get_column_letter(i + 1) for i, k in enumerate(DOC_FIELDS)}
    rng = lambda k: f"Documents!${col[k]}$2:${col[k]}${max(nd, 2)}"
    # ---- Coverage sheet
    wsc = wb.create_sheet("Coverage", 0)
    sites = {s["key"]: s for s in lib.jsonl_read(os.path.join(lib.STATE, "sites.jsonl")) if str(s.get("key", "")).startswith("official:") and s["key"].count(":") == 1}
    renames = {r[0]: r[2] for r in RENAMES_KNOWN}
    for rr in lib.jsonl_read(os.path.join(lib.STATE, "renames.jsonl")):
        renames[rr["resort_id"]] = (renames.get(rr["resort_id"], "") + f"; site redirects to {urlparse(rr['to']).netloc}").strip("; ")
    head = ["#", "Resort", "Official site status"] + [f"{t} (latest yr)" for t in DOC_TYPES] + ["Total docs", "Current docs", "Latest factsheet", "Gaps (priority 1)", "Rename notes"]
    append(wsc, head)
    for r in RESORTS:
        rid = r["resort_id"]
        row = [rid, r["name"], (sites.get(f"official:{rid}") or {}).get("status", "not crawled")]
        for t in DOC_TYPES:
            row.append(f'=IFERROR(IF(MAXIFS({rng("edition_year")},{rng("resort_id")},$A{wsc.max_row + 1},{rng("doc_type")},"{t}")=0,"",MAXIFS({rng("edition_year")},{rng("resort_id")},$A{wsc.max_row + 1},{rng("doc_type")},"{t}")),"")')
        n = wsc.max_row + 1
        row.append(f'=COUNTIFS({rng("resort_id")},$A{n})')
        row.append(f'=COUNTIFS({rng("resort_id")},$A{n},{rng("status")},"current")')
        latest_fs = next((d for d in docs.values() if rid in (d.get("resort_ids") or [d["resort_id"]]) and d["doc_type"] == "factsheet" and d["is_latest"]), None)
        if latest_fs is None:
            fs = [d for d in docs.values() if rid in (d.get("resort_ids") or [d["resort_id"]]) and d["doc_type"] == "factsheet"]
            latest_fs = max(fs, key=lambda d: d["edition_year"] or 0) if fs else None
        row.append(f'=HYPERLINK("{latest_fs["url"]}","{latest_fs["file_name"][:40]}")' if latest_fs else "")
        gaps = []
        have = {d["doc_type"] for d in docs.values() if rid in (d.get("resort_ids") or [d["resort_id"]])}
        for t in ("factsheet", "wedding", "events"):
            if t not in have:
                gaps.append(t)
        row.append(", ".join(gaps))
        row.append(renames.get(rid, ""))
        append(wsc, row)
    style_header(wsc)
    autosize(wsc, 40)
    wsc.column_dimensions["B"].width = 45
    # ---- Pages
    wsp = wb.create_sheet("Pages")
    append(wsp, ["resort_id", "resort", "url", "page_type", "title", "chars", "scraped_on", "source_type", "path"])
    for p in lib.jsonl_read(lib.PAGES_IDX):
        append(wsp, [p.get("resort_id"), BY_ID[p["resort_id"]]["name"] if p.get("resort_id") else "", p["url"], p["page_type"], p.get("title", "")[:120], p["chars"], p["scraped_on"], p.get("source_type"), p["path"]])
    style_header(wsp)
    autosize(wsp, 60)
    # ---- Sources (aggregate per domain)
    wss = wb.create_sheet("Sources")
    append(wss, ["domain", "method", "status", "pdf_found", "pdf_matched", "stored_docs", "notes", "date"])
    stored_by_source = collections.Counter(d["source"] for d in docs.values())
    for a in docs.values():
        for u in a.get("all_urls", []):
            pass
    agg = {}
    for s in lib.jsonl_read(lib.SOURCES):
        k = (s["domain"], s["method"])
        a = agg.setdefault(k, dict(s, pdf_found=0, pdf_matched=0))
        a["pdf_found"] += s.get("pdf_found", 0) or 0
        a["pdf_matched"] += s.get("pdf_matched", 0) or 0
        a["status"] = s["status"]
        a["notes"] = s.get("notes", "")
    for (dom, meth), a in sorted(agg.items(), key=lambda kv: -stored_by_source.get(kv[0][0], 0)):
        append(wss, [dom, meth, a["status"], a["pdf_found"], a["pdf_matched"], stored_by_source.get(dom, 0), a.get("notes", ""), a.get("date")])
    style_header(wss)
    autosize(wss, 50)
    # ---- Extra resorts
    wse = wb.create_sheet("Extra resorts")
    append(wse, ["url", "title", "note", "snippet", "date"])
    seen = set()
    for e in lib.jsonl_read(lib.EXTRAS):
        if e["url"] in seen:
            continue
        seen.add(e["url"])
        append(wse, [e["url"], e.get("title", ""), e.get("note", ""), e.get("snippet", "")[:200], e.get("date")])
    style_header(wse)
    autosize(wse, 60)
    # ---- Blocked
    wsb = wb.create_sheet("Blocked")
    append(wsb, ["domain", "tried", "detail", "resorts", "date"])
    blocked_sites = collections.OrderedDict()
    for s in lib.jsonl_read(os.path.join(lib.STATE, "sites.jsonl")):
        if s.get("status") in ("blocked", "blocked-partial", "unreachable", "unverified") and s.get("start_url"):
            h = urlparse(s["start_url"]).netloc
            blocked_sites.setdefault((h, s["status"]), set()).update(s.get("resort_ids") or [])
    for (h, st), ids in blocked_sites.items():
        append(wsb, [h, "GET home, sitemap, headless render", st, ", ".join(BY_ID[i]["name"] for i in sorted(ids) if i in BY_ID), lib.TODAY])
    for b in lib.jsonl_read(lib.BLOCKED):
        append(wsb, [b["domain"], b["tried"], b.get("detail", ""), "", b["date"]])
    style_header(wsb)
    autosize(wsb, 60)
    # ---- Flipbooks
    flips = [f for f in lib.jsonl_read(os.path.join(lib.STATE, "flipbooks.jsonl")) if f.get("title") != "__search__"]
    if flips:
        wsf = wb.create_sheet("Flipbooks")
        append(wsf, ["platform", "resort_id", "resort", "url", "title", "published", "download", "date"])
        for f in flips:
            append(wsf, [f["platform"], f["resort_id"], BY_ID[f["resort_id"]]["name"], f["url"], f.get("title"), f.get("published"), f.get("download"), f["date"]])
        style_header(wsf)
        autosize(wsf, 60)
    # ---- Log
    wsl = wb.create_sheet("Log")
    append(wsl, ["date", "phase", "note", "docs total"])
    lp = os.path.join(OUT, "run_log.md")
    if os.path.exists(lp):
        for line in open(lp, encoding="utf-8"):
            if line.startswith("| 20"):
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                append(wsl, cells[:4])
    style_header(wsl)
    autosize(wsl, 80)
    wb.save(os.path.join(OUT, "coverage.xlsx"))


def coverage_numbers(docs):
    have = collections.defaultdict(set)
    for d in docs.values():
        for rid in d.get("resort_ids") or [d["resort_id"]]:
            have[rid].add(d["doc_type"])
    n = lambda t: sum(1 for r in RESORTS if t in have[r["resort_id"]])
    return {"factsheet": n("factsheet"), "wedding": n("wedding"), "events": n("events"), "spa_menu": n("spa_menu"),
            "dining_menu": n("dining_menu"), "dive_prices": n("dive_prices"), "map": n("map"),
            "nothing": sum(1 for r in RESORTS if not have[r["resort_id"]]), "docs": len(docs),
            "resorts_with_docs": sum(1 for r in RESORTS if have[r["resort_id"]])}, have


def write_report(docs, extra_sections=None):
    cov, have = coverage_numbers(docs)
    stored_by_source = collections.Counter(d["source"] for d in docs.values())
    sites = {s["key"]: s for s in lib.jsonl_read(os.path.join(lib.STATE, "sites.jsonl")) if str(s.get("key", "")).startswith("official:") and s["key"].count(":") == 1}
    lines = [f"# Maldives Resort Document Library — report ({lib.TODAY})", "",
             "## Coverage", "", "| metric | resorts |", "|---|---|"]
    for k in ("factsheet", "wedding", "events", "spa_menu", "dining_menu", "dive_prices", "map", "resorts_with_docs", "nothing"):
        lines.append(f"| {k} | {cov[k]} / 182 |")
    lines += [f"| documents stored (unique by SHA-256) | {cov['docs']} |", ""]
    fs_year = {}
    for d in docs.values():
        if d["doc_type"] == "factsheet":
            for rid in d.get("resort_ids") or [d["resort_id"]]:
                fs_year[rid] = max(fs_year.get(rid, 0), d["edition_year"] or 0)
    lines += ["## Resorts without a factsheet", ""]
    lines += [f"- {r['resort_id']}. {r['name']}" for r in RESORTS if "factsheet" not in have[r["resort_id"]]] or ["- none"]
    lines += ["", "## Resorts whose newest factsheet is older than 2024", ""]
    lines += [f"- {BY_ID[rid]['resort_id']}. {BY_ID[rid]['name']} ({y or 'undated'})" for rid, y in sorted(fs_year.items()) if (y or 0) < 2024] or ["- none"]
    lines += ["", "## Missing priority-1 documents per resort (factsheet / wedding / events-MICE)", "", "| # | resort | missing | official site |", "|---|---|---|---|"]
    for r in RESORTS:
        miss = [t for t in ("factsheet", "wedding", "events") if t not in have[r["resort_id"]]]
        if miss:
            st = (sites.get("official:%d" % r["resort_id"]) or {}).get("status", "-")
            lines.append(f"| {r['resort_id']} | {r['name']} | {', '.join(miss)} | {st} |")
    lines += ["", "## Most useful sources (documents stored)", "", "| source | docs |", "|---|---|"]
    lines += [f"| {s} | {n} |" for s, n in stored_by_source.most_common(15)]
    lines += ["", "## Blocked / unreachable official sites", "", "| # | resort | status | url |", "|---|---|---|---|"]
    for r in RESORTS:
        s = sites.get(f"official:{r['resort_id']}")
        if s and s.get("status") not in ("ok", "blocked-partial"):
            lines.append(f"| {r['resort_id']} | {r['name']} | {s['status']} | {s.get('start_url', '')} |")
    lines += ["", "## Renamed resorts and aliases", "", "| # | name in list | current name / alias |", "|---|---|---|"]
    lines += [f"| {a} | {b} | {c} |" for a, b, c in RENAMES_KNOWN]
    for rr in lib.jsonl_read(os.path.join(lib.STATE, "renames.jsonl")):
        lines.append(f"| {rr['resort_id']} | {rr['name']} | official site now at {urlparse(rr['to']).netloc} (discovered) |")
    if extra_sections:
        lines += [""] + extra_sections
    lines += ["", "## Storage", "",
              "- PDFs: `library/<NNN>_<slug>/<doc_type>/<year>_<file>.pdf` (git-ignored; private). R2 upload only when `CLOUDFLARE_API_TOKEN`/`CLOUDFLARE_ACCOUNT_ID` are set.",
              "- Index: `output/documents.csv`, `output/documents.json`, `output/coverage.xlsx`.",
              "- Page text: `output/pages/<NNN>_<slug>/<type>__<host>__<path>.md`.",
              "- Fact packs: `output/factpacks/<NNN>_<slug>.json|.md`, `output/factpacks/all_resorts.csv`.",
              "- Pipeline state: `state/*.jsonl` (queue, results, sources, blocked, sites, extras); re-running any phase skips finished work.", ""]
    with open(os.path.join(OUT, "REPORT.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return cov


def build_all(extra_sections=None):
    docs = load_docs()
    mark_latest(docs)
    save_docs(docs)
    rows = write_documents(docs)
    build_workbook(docs, rows)
    cov = write_report(docs, extra_sections)
    print(json.dumps(cov))
    return cov


if __name__ == "__main__":
    build_all()
