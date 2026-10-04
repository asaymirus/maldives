"""Quality checks before finishing:
1. Re-open 10 random documents; confirm resort match and doc type from the PDF text.
2. Confirm every is_latest factsheet's edition year against its text.
3. Recalculate coverage.xlsx with LibreOffice and confirm zero formula errors.
4. List resorts without a factsheet and resorts whose newest factsheet is older than 2024.
Writes output/QA.md."""
import os, re, random, subprocess, json, shutil, collections
import lib
from lib import load_docs, pdftotext, ROOT, OUT
from resorts import RESORTS, BY_ID, match_resort, norm
from lib import classify_doc, YEAR_RE


def check_random(docs, n=10, seed=7):
    random.seed(seed)
    sample = random.sample([d for d in docs.values() if os.path.exists(os.path.join(ROOT, d["local_path"]))], n)
    rows = []
    for d in sample:
        text = pdftotext(os.path.join(ROOT, d["local_path"]), 1, 3)
        ids, review = match_resort(f"{d['file_name']} {d.get('title', '')} {text[:4000]}")
        name_ok = d["resort_id"] in ids or any(norm(a) in norm(text) for a in BY_ID[d["resort_id"]]["aliases"] if len(a) >= 5) or norm(BY_ID[d["resort_id"]]["name"]) in norm(text)
        dt, how = classify_doc(d["file_name"].split("_", 1)[-1], d.get("title", ""), text)
        rows.append({"file": d["file_name"], "resort": d["resort_name"], "resort_match_ok": bool(name_ok), "doc_type": d["doc_type"],
                     "doc_type_recheck": dt, "type_ok": dt == d["doc_type"], "snippet": re.sub(r"\s+", " ", text[:160])})
    return rows


def check_factsheet_years(docs):
    rows = []
    for d in docs.values():
        if d["doc_type"] != "factsheet" or not d["is_latest"]:
            continue
        p = os.path.join(ROOT, d["local_path"])
        if not os.path.exists(p):
            continue
        text = pdftotext(p)
        years = [int(y) for y in YEAR_RE.findall(text)]
        years = [y for y in years if y <= 2027]
        ok = (d["edition_year"] in years) if years else (d["edition_source"] in ("validity-code", "filename", "pdf-date", "last-modified"))
        rows.append({"resort": d["resort_name"], "file": d["file_name"], "edition_year": d["edition_year"], "edition_source": d["edition_source"],
                     "years_in_text": sorted(set(years))[-4:], "consistent": bool(ok)})
    return rows


def recalc_workbook():
    src = os.path.join(OUT, "coverage.xlsx")
    tmp = os.path.join(lib.STATE, "recalc")
    os.makedirs(tmp, exist_ok=True)
    try:
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, src], capture_output=True, timeout=600)
    except Exception as e:
        return {"recalculated": False, "error": str(e)}
    out = os.path.join(tmp, "coverage.xlsx")
    if not os.path.exists(out):
        return {"recalculated": False, "error": "soffice produced no file"}
    from openpyxl import load_workbook
    wb = load_workbook(out, data_only=True)
    errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:502", "Err:504", "Err:508", "Err:511"):
                    errors.append((ws.title, c.coordinate, c.value))
    cov = wb["Coverage"]
    sample = [[c.value for c in r] for r in cov.iter_rows(min_row=1, max_row=4)]
    return {"recalculated": True, "formula_errors": len(errors), "examples": errors[:10], "coverage_sample": sample}


def main():
    docs = load_docs()
    rnd = check_random(docs)
    fy = check_factsheet_years(docs)
    rc = recalc_workbook()
    have = collections.defaultdict(set)
    fs_year = {}
    for d in docs.values():
        for rid in d.get("resort_ids") or [d["resort_id"]]:
            have[rid].add(d["doc_type"])
            if d["doc_type"] == "factsheet":
                fs_year[rid] = max(fs_year.get(rid, 0), d["edition_year"] or 0)
    no_fs = [r for r in RESORTS if "factsheet" not in have[r["resort_id"]]]
    old_fs = [(BY_ID[rid], y) for rid, y in sorted(fs_year.items()) if (y or 0) < 2024]
    L = [f"# QA report ({lib.TODAY})", "", "## 1. Random re-check of 10 documents", "", "| file | resort | resort match | type stored | type re-check | ok |", "|---|---|---|---|---|---|"]
    for r in rnd:
        L.append(f"| {r['file'][:50]} | {r['resort'][:30]} | {'Y' if r['resort_match_ok'] else 'N'} | {r['doc_type']} | {r['doc_type_recheck']} | {'Y' if r['type_ok'] else 'N'} |")
    L += ["", f"Resort match confirmed: {sum(1 for r in rnd if r['resort_match_ok'])}/10; doc type stable: {sum(1 for r in rnd if r['type_ok'])}/10", ""]
    L += ["## 2. Latest factsheets: edition year vs text", "", f"{sum(1 for r in fy if r['consistent'])}/{len(fy)} consistent.", "", "| resort | file | year | source | years in text | ok |", "|---|---|---|---|---|---|"]
    for r in fy:
        if not r["consistent"]:
            L.append(f"| {r['resort'][:30]} | {r['file'][:45]} | {r['edition_year']} | {r['edition_source']} | {r['years_in_text']} | N |")
    L += ["", "## 3. Workbook recalculation", "", f"`{json.dumps({k: v for k, v in rc.items() if k != 'coverage_sample'}, default=str)[:800]}`", ""]
    if rc.get("coverage_sample"):
        L += ["Coverage sheet, first rows after recalculation:", ""] + [f"- {row[:8]}" for row in rc["coverage_sample"]]
    L += ["", f"## 4. Resorts without a factsheet ({len(no_fs)})", ""] + [f"- {r['resort_id']}. {r['name']}" for r in no_fs]
    L += ["", f"## 5. Resorts whose newest factsheet is older than 2024 ({len(old_fs)})", ""] + [f"- {r['resort_id']}. {r['name']} ({y or 'undated'})" for r, y in old_fs]
    with open(os.path.join(OUT, "QA.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")
    print("\n".join(L[:40]))
    return rc


if __name__ == "__main__":
    main()
